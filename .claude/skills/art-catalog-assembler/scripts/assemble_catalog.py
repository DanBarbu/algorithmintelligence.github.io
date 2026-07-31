#!/usr/bin/env python3
"""Replace each placeholder image in the catalog .docx with its matched
artefact image, in place, preserving all surrounding text/formatting.

Reads placeholders.json (from extract_placeholders.py) and mapping.csv
(from build_mapping.py or match_by_text.py, optionally hand-edited) and:

  1. For each mapped row, swaps the placeholder picture's blip target for
     a newly-added image part, resizing to fit within a max width/height
     box while preserving the image's own aspect ratio.
  2. Strips the "Figure Placeholder: " marker prefix from the caption
     paragraph, leaving a clean caption.
  3. If the row has a "hyperlink" column value, inserts a new paragraph
     right after the caption with a clickable "Source: <url>" link --
     provenance for artefacts whose photo came from an external listing.
  4. Leaves any UNMATCHED_PLACEHOLDER rows untouched (still shows the
     grey placeholder box) so gaps stay visible rather than silently
     disappearing.

Usage:
    python assemble_catalog.py catalog.docx placeholders.json mapping.csv \
        images_dir out.docx [--marker "Figure Placeholder:"] \
        [--max-width-in 6.0] [--max-height-in 8.0]
"""
import argparse
import csv
import json
import os

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Pt

EMU_PER_IN = 914400


def set_extent(drawing_el, cx, cy):
    # wp:extent (controls the visible size in Word) and the a:ext inside
    # pic:spPr/a:xfrm (the picture's own frame) must both be updated, or
    # Word keeps rendering the old placeholder box size.
    extent_el = drawing_el.find(".//" + qn("wp:extent"))
    if extent_el is not None:
        extent_el.set("cx", str(cx))
        extent_el.set("cy", str(cy))
    for ext_el in drawing_el.findall(".//" + qn("a:ext")):
        ext_el.set("cx", str(cx))
        ext_el.set("cy", str(cy))


def replace_picture(doc, picture_paragraph_index, image_path, max_w_emu, max_h_emu):
    paragraph = doc.paragraphs[picture_paragraph_index]
    drawing_el = paragraph._p.find(".//" + qn("w:drawing"))
    if drawing_el is None:
        raise ValueError(f"No drawing found in paragraph {picture_paragraph_index}")

    blip_el = drawing_el.find(".//" + qn("a:blip"))
    if blip_el is None:
        raise ValueError(f"No blip found in paragraph {picture_paragraph_index}")

    rId, image = doc.part.get_or_add_image(image_path)
    blip_el.set(qn("r:embed"), rId)

    px_w, px_h = image.px_width, image.px_height
    scale = min(max_w_emu / px_w, max_h_emu / px_h)
    cx, cy = int(px_w * scale), int(px_h * scale)
    set_extent(drawing_el, cx, cy)


def strip_marker(doc, caption_paragraph_index, marker, description):
    paragraph = doc.paragraphs[caption_paragraph_index]
    if not paragraph.runs:
        return
    paragraph.runs[0].text = description
    for r in paragraph.runs[1:]:
        r.text = ""


def add_hyperlink_run(paragraph, url, text):
    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rPr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "16")  # 8pt
    rPr.append(sz)
    run.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def insert_source_line(doc, caption_paragraph_index, url):
    # Insert as a new paragraph anchored immediately before the paragraph
    # that currently follows the caption -- i.e. right after the caption.
    # Must be called in descending caption_paragraph_index order across a
    # batch so indices computed before any insertion stay valid (inserting
    # a paragraph shifts every index after it).
    caption_paragraph = doc.paragraphs[caption_paragraph_index]
    anchor = doc.paragraphs[caption_paragraph_index + 1]
    new_para = anchor.insert_paragraph_before()
    new_para.alignment = caption_paragraph.alignment
    prefix = new_para.add_run("Source: ")
    prefix.italic = True
    prefix.font.size = Pt(8)
    add_hyperlink_run(new_para, url, url)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("docx_path")
    ap.add_argument("placeholders_json")
    ap.add_argument("mapping_csv")
    ap.add_argument("images_dir")
    ap.add_argument("out_docx")
    ap.add_argument("--marker", default="Figure Placeholder:")
    ap.add_argument("--max-width-in", type=float, default=6.0)
    ap.add_argument("--max-height-in", type=float, default=8.0)
    args = ap.parse_args()

    placeholders = {p["seq"]: p for p in json.load(open(args.placeholders_json, encoding="utf-8"))}
    max_w_emu = int(args.max_width_in * EMU_PER_IN)
    max_h_emu = int(args.max_height_in * EMU_PER_IN)

    doc = Document(args.docx_path)

    skip_flags = ("UNMATCHED_PLACEHOLDER", "UNUSED_IMAGE", "UNUSED_SLIDE")
    rows = [r for r in csv.DictReader(open(args.mapping_csv, encoding="utf-8"))
            if r["flag"] not in skip_flags and r["image_file"]]
    n_skipped = sum(1 for r in csv.DictReader(open(args.mapping_csv, encoding="utf-8"))
                    if r["flag"] == "UNMATCHED_PLACEHOLDER")

    for row in rows:
        seq = int(row["seq"])
        ph = placeholders[seq]
        image_path = os.path.join(args.images_dir, row["image_file"])
        replace_picture(doc, ph["picture_paragraph_index"], image_path, max_w_emu, max_h_emu)
        strip_marker(doc, ph["caption_paragraph_index"], args.marker, ph["description"])

    # Hyperlink insertion changes paragraph count, so it must run as a
    # second pass, in descending index order, after all the (count-neutral)
    # image/caption edits above are done.
    n_links = 0
    hyperlink_rows = sorted(
        (r for r in rows if r.get("hyperlink")),
        key=lambda r: placeholders[int(r["seq"])]["caption_paragraph_index"],
        reverse=True,
    )
    for row in hyperlink_rows:
        ph = placeholders[int(row["seq"])]
        insert_source_line(doc, ph["caption_paragraph_index"], row["hyperlink"])
        n_links += 1

    doc.save(args.out_docx)
    print(f"Assembled {len(rows)} images ({n_links} with a source hyperlink) -> {args.out_docx} "
          f"({n_skipped} placeholders left unmatched)")


if __name__ == "__main__":
    main()
